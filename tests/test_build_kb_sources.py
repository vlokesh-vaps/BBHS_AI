"""Tests for building a vector source from PDF and DOCX files."""

from pathlib import Path

from docx import Document as DocxDocument
import pymupdf

from scripts.build_kb import load_docx_documents, load_source_documents


def test_load_docx_documents_extracts_text_and_tables(tmp_path: Path):
    source = tmp_path / "school-notice.docx"
    doc = DocxDocument()
    doc.add_heading("School Visit Information", level=1)
    doc.add_paragraph("Please contact Admissions before visiting the campus.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Contact"
    table.cell(0, 1).text = "Email"
    table.cell(1, 0).text = "Admissions"
    table.cell(1, 1).text = "admission@example.edu.in"
    doc.save(source)

    documents = load_docx_documents(source)

    assert len(documents) == 1
    assert "# School Visit Information" in documents[0].page_content
    assert "Please contact Admissions" in documents[0].page_content
    assert "admission@example.edu.in" in documents[0].page_content
    assert documents[0].metadata["file_type"] == "docx"


def test_load_source_documents_scans_pdf_and_docx(tmp_path: Path):
    pdf = tmp_path / "school.pdf"
    pdf_doc = pymupdf.open()
    page = pdf_doc.new_page()
    page.insert_text((50, 50), "PDF school policy content")
    pdf_doc.save(pdf)
    pdf_doc.close()
    docx = tmp_path / "school.docx"
    doc = DocxDocument()
    doc.add_paragraph("DOCX school policy content")
    doc.save(docx)

    documents = load_source_documents(tmp_path)
    combined = "\n".join(document.page_content for document in documents)

    assert "PDF school policy content" in combined
    assert "DOCX school policy content" in combined
    assert {document.metadata["file_type"] for document in documents} == {"pdf", "docx"}
